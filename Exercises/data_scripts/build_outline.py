"""Fill the QLS-MiCM outline template (docx) and write Outline.md from one source.

Run from the repo root:  python Exercises/data_scripts/build_outline.py <path/to/MiCM - Workshop Outline Template.docx>
"""
import copy
import sys
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[2]

TITLE = "Data Visualization in Python"
LEAD = "Zahra Tavakol Hamedani (Instructor) · Sameena Karsan (Facilitator)"
CONTEXT = "Code to Manuscript Bootcamp – Day 1 · Monday, October 26, 2026, 2:30–4:30 pm · McIntyre Room 1027"
DURATION = "2 hours"
PREREQS = [
    "Basic Python: variables, lists/dictionaries, loops and calling functions.",
    "Basic pandas: loading a CSV into a DataFrame and selecting rows/columns.",
    "A Google account to run the notebook in Google Colab (preferred), or a local Python 3 installation "
    "(≥ 3.10) with Jupyter, pandas, matplotlib, seaborn, scikit-learn, umap-learn and adjustText.",
]
SUMMARY = (
    "This hands-on session explores the Python data visualization ecosystem, focusing on the relationship between "
    "Matplotlib Figure/Axes objects and Seaborn's high-level statistical plotting functions. Using a public bulk "
    "RNA-seq dataset (SARS-CoV-2-infected human ocular tissue, GEO GSE164073; the same dataset used on Day 2 in R), "
    "participants create quality-control plots, distributions, volcano/MA plots, heatmaps and PCA/t-SNE/UMAP "
    "embeddings. They then assemble a multi-panel figure at journal dimensions with a consistent, colour-blind-safe "
    "style and export it as vector (PDF/SVG) and high-resolution raster files for Inkscape (Day 4) and LaTeX (Day 5)."
)
OBJECTIVES = [
    "Explain the Matplotlib Figure/Axes model and use the object-oriented interface to build and customise plots.",
    "Reshape data into tidy (long) format and choose between axes-level and figure-level Seaborn functions.",
    "Create common biological visualizations: distributions, strip/box plots, volcano and MA plots, "
    "(clustered) heatmaps, and PCA/t-SNE/UMAP embeddings.",
    "Assemble a multi-panel figure at journal column width with consistent colours, fonts and panel labels.",
    "Export figures as editable vector (PDF/SVG) and ≥300 dpi raster files ready for post-processing.",
]
# (module title, minutes, [(section, [subsections])], exercise)
CONTENT = [
    ("Module 0 – Setup & the dataset", 10, [
        ("Opening the notebook in Colab; installing umap-learn and adjustText", []),
        ("The GSE164073 dataset: 3 tissues × 2 conditions × 3 donors", ["Counts, normalised expression, DE results"]),
        ("Defining a project-wide colour-blind-safe palette (Okabe–Ito)", []),
    ], None),
    ("Module 1 – Anatomy of a Matplotlib figure", 15, [
        ("Figure vs Axes; implicit (plt.) vs explicit (fig, ax) interfaces", []),
        ("Bar chart of library sizes; labels, units, colours and custom legends", []),
    ], "Exercise 1 – horizontal bar chart coloured by condition with a median line (5 min)"),
    ("Module 2 – Seaborn, tidy data & distributions", 20, [
        ("Wide vs long (tidy) data with melt/merge", []),
        ("Axes-level vs figure-level functions", ["Per-sample box plots (QC)", "Per-sample density curves faceted by tissue (optional)"]),
        ("Showing individual points for small n: bar of means vs box + strip plot", []),
    ], "Exercise 2 – small multiples of four inflammatory genes with catplot (5 min)"),
    ("Module 3 – Relationships: scatter, MA & volcano plots", 20, [
        ("Volcano plot: thresholds, colour by DE status, non-overlapping gene labels (adjustText)", []),
        ("Rasterising dense point layers inside vector figures", []),
    ], "Exercise 3 – MA plot (5 min)"),
    ("Module 4 – Heatmaps & clustered heatmaps", 15, [
        ("Choosing colour maps: sequential vs diverging; avoiding rainbow", []),
        ("Sample–sample correlation heatmap (optional)", []),
        ("Clustered heatmap of variable genes with z-scores and annotation colour bars", []),
        ("Pitfall: heatmaps of DE-selected genes always separate the groups (circular analysis)", []),
    ], "Exercise 4 – clustered heatmap of the top 30 DE genes in cornea (5 min)"),
    ("Module 5 – Dimensionality reduction: PCA, t-SNE, UMAP", 15, [
        ("Intuition only: what each method preserves and how to read the axes", []),
        ("Side-by-side embeddings using hue (tissue) and style (condition)", []),
        ("Within-tissue PCA: the input decides which effect a PCA can show", []),
    ], "Exercise 5 (optional / live demo) – sensitivity of t-SNE to perplexity; what to report in Methods (3 min)"),
    ("Module 6 – Multi-panel figures & high-resolution export", 20, [
        ("Journal sizing in mm, paper-style rcParams, subplot_mosaic layouts and panel labels", []),
        ("Combining a clustered heatmap inside a layout (SciPy ordering + sns.heatmap)", []),
        ("Exporting PDF/SVG with editable text and a 600 dpi PNG; checking size and fonts", []),
        ("Downloading figures from Colab for Day 4", []),
    ], "Exercise 6 – build and export a 3-panel figure for Day 4 (8 min; can be finished after class)"),
    ("Wrap-up – publication-figure checklist, next sessions & feedback", 5, [], None),
]


def write_md():
    L = [f"# {TITLE}", "", f"**Workshop Lead:** {LEAD}  ", f"**Context:** {CONTEXT}  ",
         "**Registration link:** NA  ", f"**Approximate duration:** {DURATION}", "", "## Prerequisites", ""]
    L += [f"- {p}" for p in PREREQS]
    L += ["", "## Summary", "", SUMMARY, "", "## Learning Objectives", ""]
    L += [f"{i}. {o}" for i, o in enumerate(OBJECTIVES, 1)]
    L += ["", "## Content", ""]
    for i, (m, t, secs, ex) in enumerate(CONTENT, 1):
        L.append(f"{i}. **{m} ({t} minutes)**")
        for s, subs in secs:
            L.append(f"   1. {s}")
            L += [f"      1. {x}" for x in subs]
        if ex:
            L.append(f"   1. **{ex}**")
    (ROOT / "Outline/Outline.md").write_text("\n".join(L) + "\n")


def write_docx(template):
    d = Document(template)
    P = d.paragraphs

    def clone(src, anchor, txt, bold=None):
        new = copy.deepcopy(src._p)
        for r in new.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r")[1:]:
            new.remove(r)
        anchor._p.addprevious(new)
        from docx.text.paragraph import Paragraph
        p = Paragraph(new, src._parent)
        if p.runs:
            p.runs[0].text = txt
            if bold is not None: p.runs[0].bold = bold
        else:
            r = p.add_run(txt); r.bold = bold
        return p

    # header fields
    P[1].runs[0].text = TITLE; P[1].runs[1].text = ""; P[1].runs[2].text = ""
    P[2].runs[1].text = ": "; r = P[2].add_run(LEAD); r.bold = False
    r = P[2].add_run("\n" + CONTEXT); r.bold = False; r.italic = True
    P[4].runs[0].text = f"Approximate duration: {DURATION}"
    # prerequisites (paragraphs 6-8)
    pre_src, pre_anchor = P[6], P[9]
    for t in PREREQS:
        clone(pre_src, pre_anchor, t, False)
    # summary
    for r in P[9].runs[2:]: r.text = ""
    P[9].runs[1].text = ": "
    r = P[9].add_run(SUMMARY); r.bold = False
    # objectives (12-16)
    obj_src, obj_anchor = P[12], P[17]
    for t in OBJECTIVES:
        clone(obj_src, obj_anchor, t, False)
    for r in P[11].runs[2:]: r.text = ""
    # content (18-29)
    lvl0, lvl1, lvl2 = P[18], P[19], P[20]
    end = d.add_paragraph()
    for m, t, secs, ex in CONTENT:
        clone(lvl0, end, f"{m} ({t} min)", True)
        for s, subs in secs:
            clone(lvl1, end, s, False)
            for x in subs:
                clone(lvl2, end, x, False)
        if ex:
            clone(lvl1, end, "Hands-on: " + ex, True)
    # remove template placeholders
    for p in [P[i] for i in (6, 7, 8, 12, 13, 14, 15, 16)] + P[18:30] + [end]:
        p._p.getparent().remove(p._p)
    out = ROOT / "Outline/MiCM - Data Visualization in Python Workshop Outline.docx"
    d.save(out)
    return out


if __name__ == "__main__":
    write_md()
    if len(sys.argv) > 1:
        print(write_docx(sys.argv[1]))
